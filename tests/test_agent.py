"""Unit and integration tests for DocAgent core planning, tools, reflection, and document building."""
import os
import pytest
from docx import Document

from app.models.schemas import AgentRequest, TaskStatus
from app.agent.core import DocAgent
from app.agent.planner import Planner
from app.agent.tools import tool_registry
from app.agent.reflection import ReflectionReviewer
from app.document.builder import DocxBuilder
from app.llm.client import FallbackLLMClient

def test_planner_creates_tasks_and_assumptions():
    llm = FallbackLLMClient()
    planner = Planner(llm)
    
    # Test ambiguous request
    request = "Migrate legacy monolith to microservices next month, keep cloud costs under $50k, CTO wants k8s, architect wants serverless."
    plan = planner.create_plan(request)
    
    assert len(plan.tasks) >= 5
    assert len(plan.assumptions_made) >= 2
    assert any("50,000" in a or "budget" in a.lower() for a in plan.assumptions_made)
    assert any("kubernetes" in a.lower() or "serverless" in a.lower() for a in plan.assumptions_made)
    assert plan.tasks[0].status == TaskStatus.PENDING

def test_tools_execution():
    res = tool_registry.research_tool("Cloud Migration")
    assert res["status"] == "success"
    assert len(res["findings"]) > 0

    metrics = tool_registry.metrics_calculator_tool([{"amount": 10000}, {"amount": 25000}])
    assert metrics["calculated_total"] == 35000

def test_docx_builder_creates_valid_file(tmp_path):
    builder = DocxBuilder()
    builder.add_title_block(
        title="Unit Test Specification",
        doc_type="Technical Spec",
        audience="Engineering",
        assumptions=["Test assumption 1", "Test assumption 2"]
    )
    from app.models.schemas import DocumentSection, TableData
    sec = DocumentSection(
        title="Test Heading 1",
        level=1,
        content="This is a test paragraph verifying python-docx rendering.",
        callout="Executive note: test passed",
        table=TableData(
            headers=["Service", "Status"],
            rows=[["Auth", "OK"], ["Billing", "OK"]],
            caption="Table 1: Test"
        )
    )
    builder.add_section(sec)
    
    meta = builder.save("test_document")
    assert os.path.exists(meta.file_path)
    assert meta.file_size_bytes > 0
    
    # Verify file is readable as docx
    loaded_doc = Document(meta.file_path)
    assert len(loaded_doc.paragraphs) > 0
    assert len(loaded_doc.tables) >= 2  # metadata table + custom section table

def test_doc_agent_end_to_end_standard():
    agent = DocAgent(provider_override="fallback")
    req = AgentRequest(request="Create a formal Project Proposal for customer onboarding automation at Acme Corp.")
    res = agent.run(req)
    
    assert res.success is True
    assert len(res.plan.tasks) >= 5
    assert res.reflection_report.score >= 85
    assert os.path.exists(res.document.file_path)
    assert res.document.download_url.startswith("/download/")
